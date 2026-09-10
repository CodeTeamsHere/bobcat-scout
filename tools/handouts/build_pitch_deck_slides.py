# Add three honest-pitch slides to the Bobcat Scout deck by cloning the
# existing designed slides and rewriting their text, so the new slides
# inherit the deck's exact look.
import copy, io, sys
from pptx import Presentation

SRC = r'C:\Users\kishg\Downloads\Bobcat_Scout_pitch_deck.pptx'
OUT = r'C:\Users\kishg\Downloads\Bobcat_Scout_pitch_deck_v2.pptx'

prs = Presentation(SRC)


def clone_slide(prs, index):
    """Deep-copy slide `index` onto a new slide at the end of the deck."""
    src = prs.slides[index]
    dst = prs.slides.add_slide(src.slide_layout)
    # drop placeholders the layout added
    for shp in list(dst.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(shp._element))
    # carry the slide's own background across; add_slide() only gives us the layout's
    ns = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
    src_bg = src._element.find(ns + 'cSld/' + ns + 'bg')
    if src_bg is not None:
        dst_cSld = dst._element.find(ns + 'cSld')
        old_bg = dst_cSld.find(ns + 'bg')
        if old_bg is not None:
            dst_cSld.remove(old_bg)
        dst_cSld.insert(0, copy.deepcopy(src_bg))
    return dst


def move_slide(prs, from_idx, to_idx):
    ids = prs.slides._sldIdLst
    items = list(ids)
    ids.remove(items[from_idx])
    ids.insert(to_idx, items[from_idx])


def set_text(shape, text):
    """Replace a shape's text, keeping the formatting of its first run.

    Runs must sit before <a:endParaRPr> or PowerPoint silently drops the text,
    so build each paragraph through python-pptx's own add_run().
    """
    from pptx.text.text import _Paragraph
    tf = shape.text_frame
    lines = text.split('\n')
    p0 = tf.paragraphs[0]
    donor_rPr = None
    if p0.runs:
        rPr = p0.runs[0]._r.find(
            '{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
        if rPr is not None:
            donor_rPr = copy.deepcopy(rPr)

    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)

    def write(par, s):
        run = par.add_run()
        if donor_rPr is not None:
            old = run._r.find(
                '{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
            if old is not None:
                run._r.remove(old)
            run._r.insert(0, copy.deepcopy(donor_rPr))
        run.text = s

    write(p0, lines[0])
    template = copy.deepcopy(p0._p)
    for line in lines[1:]:
        np = copy.deepcopy(template)
        tf._txBody.append(np)
        par = _Paragraph(np, tf)
        for r in list(par.runs):
            r._r.getparent().remove(r._r)
        write(par, line)


def fill(slide, mapping):
    """mapping: {shape_name: new_text}"""
    for shp in slide.shapes:
        if shp.name in mapping and shp.has_text_frame:
            set_text(shp, mapping[shp.name])


# ---------------------------------------------------------------- slide A
# clone of slide 4 (two-column compare) -> "IT DOES NOT REPLACE WHAT WE DO"
a = clone_slide(prs, 3)
fill(a, {
    'Text 0': 'IT DOES NOT REPLACE WHAT WE DO',
    'Text 1': 'Run it beside what we already do.',
    'Text 3': 'DURING THE EVENT',
    'Text 4': 'Both',
    'Text 5': 'systems at once',
    'Text 7': 'Normal scouting keeps going exactly as it is\nTwo or three people also run Bobcat Scout\nSame matches, two independent records',
    'Text 9': 'AFTER THE EVENT',
    'Text 10': 'One',
    'Text 11': 'honest answer',
    'Text 13': 'Which one got the row in faster\nWhich one matched the official scores\nWhich one the scouters wanted to keep',
    'Text 14': 'If it does not win on the numbers, we stop. Our existing data was never interrupted.',
})

# ---------------------------------------------------------------- slide B
# clone of slide 6 (2x2 numbered grid) -> "WHAT COULD GO WRONG"
b = clone_slide(prs, 5)
fill(b, {
    'Text 0': 'WHAT COULD GO WRONG',
    'Text 1': 'The honest list, with answers.',
    'Text 3': '01',
    'Text 4': 'Loud stands',
    'Text 5': 'The one real risk. Close-talk headsets fix it, and the pilot measures it.',
    'Text 7': '02',
    'Text 8': 'Nobody wants to talk',
    'Text 9': 'Then they type. Both paths write identical rows to the sheet.',
    'Text 11': '03',
    'Text 12': 'It mishears a word',
    'Text 13': 'Nothing sends until a human sees the filled form and confirms it.',
    'Text 15': '04',
    'Text 16': 'One student built it',
    'Text 17': 'No server, no code to edit, and two more members trained before the season.',
})

# ---------------------------------------------------------------- slide C
# clone of slide 7 (three big-word cards) -> "THE ASK"
c = clone_slide(prs, 6)
fill(c, {
    'Text 0': 'THE ASK',
    'Text 1': 'One event. Nothing turned off.',
    'Text 2': 'One',
    'Text 3': 'Event',
    'Text 4': 'Run it beside our normal scouting for a single competition.',
    'Text 5': 'Three',
    'Text 6': 'Scouters',
    'Text 7': 'Volunteers who will say honestly if it turns out worse.',
    'Text 8': 'Zero',
    'Text 9': 'Risk',
    'Text 10': 'Normal scouting never stops, so nothing is lost if we say no.',
})

# widen the narrow boxes the cloned text now overflows
from pptx.util import Inches
for name, w in (('Text 5', Inches(3.0)), ('Text 11', Inches(3.0))):
    for shp in a.shapes:
        if shp.name == name:
            shp.width = w
for name in ('Text 3', 'Text 7', 'Text 11', 'Text 15'):
    for shp in b.shapes:
        if shp.name == name:
            shp.width = Inches(1.2)

# put the three new slides just before the closing slide (index 7)
n = len(prs.slides._sldIdLst)          # 11 after cloning
move_slide(prs, n - 3, 7)              # slide A -> position 8
move_slide(prs, n - 2, 8)              # slide B -> position 9
move_slide(prs, n - 1, 9)              # slide C -> position 10

prs.save(OUT)
print('saved', OUT, 'slides =', len(prs.slides._sldIdLst))
